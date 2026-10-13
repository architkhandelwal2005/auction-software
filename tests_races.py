"""Two changes at once: a double tap, the stage and a phone both in use, a
slow phone whose request lands late. Each test fires its requests together
from threads, with every database read slowed down so any gap between a check
and the write that depends on it is wide open, then checks the auction still
adds up: every team's purse equals its budget less what it paid.
"""
import os, sys, time, json, threading, sqlite3, tempfile
os.environ['DATABASE_URL'] = ''
os.environ['SUPABASE_URL'] = ''
os.environ['SUPABASE_SERVICE_KEY'] = ''
# RACE_APP_DIR: run these against another copy of the app (an older one, to
# see which races it loses).
APP_DIR = os.environ.get('RACE_APP_DIR', r'D:\auction software')
sys.path.insert(0, APP_DIR)
os.chdir(APP_DIR)
import app as A

def check(label, cond, detail=''):
    print(('  PASS  ' if cond else '  FAIL  ') + label + (' :: ' + detail if detail else ''))
    if not cond: check.failed += 1
check.failed = 0

# Backups go to a scratch folder, not over the real ones beside app.py.
BACKUPS = tempfile.mkdtemp(prefix='race_backups_')
A.excel_backup_path = lambda auction_id: os.path.join(BACKUPS, 'auction_backup_%s.xlsx' % auction_id)

# Every read made while handling a request in a worker thread waits a moment.
SLOW = {'on': False, 'delay': 0.12}
MAIN = threading.main_thread()
class SlowCursor(sqlite3.Cursor):
    def execute(self, sql, params=()):
        r = super().execute(sql, params)
        if SLOW['on'] and threading.current_thread() is not MAIN and sql.lstrip().upper().startswith('SELECT'):
            time.sleep(SLOW['delay'])
        return r
class SlowConnection(sqlite3.Connection):
    def cursor(self, factory=SlowCursor):
        return super().cursor(factory)
    def execute(self, sql, params=()):
        return self.cursor().execute(sql, params)
_connect = sqlite3.connect
def slow_connect(*a, **k):
    k.setdefault('factory', SlowConnection)
    return _connect(*a, **k)
sqlite3.connect = slow_connect

PW = os.environ.get('ADMIN_PASSWORD', 'admin@123')
def admin_client():
    c = A.app.test_client()
    c.post('/api/auth/login', json={'role': 'admin', 'password': PW})
    return c

def together(*calls):
    """Run the calls at the same moment, each on its own client; their
    responses, in order."""
    out = [None] * len(calls)
    start = threading.Barrier(len(calls))
    def run(i, fn):
        start.wait()
        out[i] = fn()
    threads = [threading.Thread(target=run, args=(i, fn)) for i, fn in enumerate(calls)]
    SLOW['on'] = True
    for t in threads: t.start()
    for t in threads: t.join()
    SLOW['on'] = False
    return out

admin = admin_client()
def new_auction(name, teams=(('Reds', 100), ('Blues', 100)), players=(('P1', 1),)):
    aid = admin.post('/api/auctions', json={'name': name}).get_json()['auction']['id']
    admin.post('/api/auctions/%d/open' % aid)
    tids = [admin.post('/api/teams', json={'name': n, 'total_budget': b}).get_json()['id'] for n, b in teams]
    pids = [admin.post('/api/players', json={'name': n, 'category': 'A', 'base_price': bp}).get_json()['id']
            for n, bp in players]
    return aid, tids, pids

def on(aid):
    """A client working on auction aid (its own session)."""
    c = admin_client()
    c.post('/api/auctions/%d/open' % aid)
    return c

def books_balance(label):
    teams = admin.get('/api/teams').get_json()
    players = admin.get('/api/players').get_json()
    bad = []
    for t in teams:
        paid = sum(p['sold_price'] or 0 for p in players if p['status'] == 'sold' and p['team_id'] == t['id'])
        if abs(t['total_budget'] - paid - t['remaining_budget']) > 1e-6:
            bad.append('%s: budget %s, paid %s, purse %s' % (t['name'], t['total_budget'], paid, t['remaining_budget']))
        if t['remaining_budget'] < -1e-6:
            bad.append('%s purse below zero' % t['name'])
    for p in players:
        if (p['status'] == 'sold') != (p['team_id'] is not None):
            bad.append('%s is %s with team %s' % (p['name'], p['status'], p['team_id']))
    check(label + ': the books balance', not bad, '; '.join(bad))

def status(pid):
    return next(p for p in admin.get('/api/players').get_json() if p['id'] == pid)

def stage():
    return admin.get('/api/live_state').get_json()['auction_state']

print('Sales')
aid, (red, blue), (p1,) = new_auction('Double sale')
a, b = on(aid), on(aid)
r = together(lambda: a.post('/api/sell_player', json={'player_id': p1, 'team_id': red, 'sold_price': 5}),
             lambda: b.post('/api/sell_player', json={'player_id': p1, 'team_id': blue, 'sold_price': 6}))
check('one player sold twice at once: one sale stands', sorted(x.status_code for x in r) == [200, 409],
      str([x.status_code for x in r]))
books_balance('double sale')

aid, (red,), (p1, p2) = new_auction('One purse, two sales', teams=(('Reds', 10),), players=(('P1', 1), ('P2', 1)))
a, b = on(aid), on(aid)
r = together(lambda: a.post('/api/sell_player', json={'player_id': p1, 'team_id': red, 'sold_price': 8}),
             lambda: b.post('/api/sell_player', json={'player_id': p2, 'team_id': red, 'sold_price': 8}))
check('two sales one purse can cover only once: one stands', sorted(x.status_code for x in r) == [200, 400],
      str([x.status_code for x in r]))
books_balance('one purse, two sales')

aid, (red,), (p1, p2) = new_auction('Sale and allotment', teams=(('Reds', 10),), players=(('P1', 1), ('P2', 1)))
a, b = on(aid), on(aid)
r = together(lambda: a.post('/api/sell_player', json={'player_id': p1, 'team_id': red, 'sold_price': 7}),
             lambda: b.post('/api/players/edit_sale', json={'player_id': p2, 'team_id': red, 'sold_price': 7, 'status': 'sold'}))
check('a sale and an allotment the purse covers only once: one stands', sorted(x.status_code for x in r) == [200, 400],
      str([x.status_code for x in r]))
books_balance('sale and allotment')

print('Undo')
aid, (red,), (p1, p2) = new_auction('Double undo', teams=(('Reds', 100),), players=(('P1', 1), ('P2', 1)))
admin.post('/api/sell_player', json={'player_id': p1, 'team_id': red, 'sold_price': 5})
admin.post('/api/sell_player', json={'player_id': p2, 'team_id': red, 'sold_price': 7})
a, b = on(aid), on(aid)
r = together(lambda: a.post('/api/undo'), lambda: b.post('/api/undo'))
names = sorted((x.get_json() or {}).get('player_name') or '' for x in r)
check('two undos at once undo two different sales', names == ['P1', 'P2'], str(names))
books_balance('double undo')

aid, (red,), (p1, p2) = new_auction('Undo during a sale', teams=(('Reds', 100),), players=(('P1', 1), ('P2', 1)))
admin.post('/api/sell_player', json={'player_id': p1, 'team_id': red, 'sold_price': 5})
a, b = on(aid), on(aid)
together(lambda: a.post('/api/undo'),
         lambda: b.post('/api/sell_player', json={'player_id': p2, 'team_id': red, 'sold_price': 9}))
books_balance('undo during a sale')

print('Pass and bring back')
aid, (red,), (p1,) = new_auction('Pass during a sale', teams=(('Reds', 100),))
a, b = on(aid), on(aid)
r = together(lambda: a.post('/api/sell_player', json={'player_id': p1, 'team_id': red, 'sold_price': 5}),
             lambda: b.post('/api/action/pass', json={'player_id': p1}))
books_balance('pass during a sale')
admin.post('/api/sell_player', json={'player_id': p1, 'team_id': red, 'sold_price': 5})   # sold either way now
check('a late pass of a sold player is refused', admin.post('/api/action/pass', json={'player_id': p1}).status_code == 409)
check('and the player stays sold', status(p1)['status'] == 'sold')
r = admin.post('/api/action/revive', json={'player_id': p1})
check('bringing back a sold player does nothing', r.get_json()['count'] == 0 and status(p1)['status'] == 'sold')
books_balance('late pass and revive')

aid, (red,), (p1,) = new_auction('Double pass', teams=(('Reds', 100),), players=(('P1', 4),))
a, b = on(aid), on(aid)
r = together(lambda: a.post('/api/action/pass', json={'player_id': p1}),
             lambda: b.post('/api/action/pass', json={'player_id': p1}))
check('a double-tapped pass succeeds once', [x.status_code for x in r] == [200, 200]
      and sorted(bool(x.get_json().get('already')) for x in r) == [False, True], str([x.get_json() for x in r]))
r = together(lambda: a.post('/api/action/revive', json={'player_id': p1, 'half_price': True}),
             lambda: b.post('/api/action/revive', json={'player_id': p1, 'half_price': True}))
check('a double-tapped half-price round halves once', status(p1)['base_price'] == 2, str(status(p1)['base_price']))
check('one undo then reverses the pass', admin.post('/api/undo').get_json().get('success')
      and admin.post('/api/undo').get_json().get('success') is False)

print('Bids')
aid, (red, blue), (p1, p2) = new_auction('Late bids', players=(('P1', 1), ('P2', 1)))
admin.post('/api/auction/state', json={'current_player': 'P1', 'current_player_id': str(p1), 'current_bid': 1})
ok = admin.post('/api/auction/state', json={'for_player_id': str(p1), 'current_bid': 2, 'bidder_team_id': str(red)})
check('a bid for the player on the block goes through', ok.status_code == 200 and stage()['current_bid'] == '2')
check('and the player it names is not stored', 'for_player_id' not in stage())
admin.post('/api/sell_player', json={'player_id': p1, 'team_id': red, 'sold_price': 2})
admin.post('/api/auction/state', json={'current_player': 'P2', 'current_player_id': str(p2), 'current_bid': 1})
late = admin.post('/api/auction/state', json={'for_player_id': str(p1), 'current_bid': 9, 'bidder_team_id': str(blue)})
check('a late bid for the sold player is refused', late.status_code == 409, str(late.status_code))
check('and the next player keeps their own bid', stage()['current_bid'] == '1' and not stage().get('bidder_team_id'),
      str(stage().get('current_bid')))
back = admin.post('/api/auction/state', json={'current_player': 'P1', 'current_player_id': str(p1), 'current_bid': 2})
check('a late request cannot put a sold player back on the block', back.status_code == 409
      and stage()['current_player'] == 'P2')
admin.post('/api/auction/state', json={'current_player': '', 'current_bid': 0})
empty = admin.post('/api/auction/state', json={'for_player_id': str(p2), 'current_bid': 3})
check('a bid with nobody on the block is refused', empty.status_code == 409)

admin.post('/api/auction/state', json={'current_player': 'P2', 'current_player_id': str(p2), 'current_bid': 1})
screens = [on(aid) for _ in range(8)]
teams = {red: 'Reds', blue: 'Blues'}
def bid(i):
    tid = (red, blue)[i % 2]
    body = {'for_player_id': str(p2), 'current_bid': 10 + i, 'bidder_team_id': str(tid), 'bidder_team_name': teams[tid]}
    if i % 2:   # other screens send the fields in another order
        body = dict(reversed(list(body.items())))
    return screens[i].post('/api/auction/state', json=body)
r = together(*[lambda i=i: bid(i) for i in range(8)])
check('eight screens bidding at once: none fails', all(x.status_code == 200 for x in r), str([x.status_code for x in r]))
st = stage()
i = int(st['current_bid']) - 10
check('the stage shows one whole bid, not halves of two', st['bidder_team_id'] == str((red, blue)[i % 2])
      and st['bidder_team_name'] == teams[(red, blue)[i % 2]], json.dumps({k: st.get(k) for k in ('current_bid', 'bidder_team_id', 'bidder_team_name')}))

print('End Auction')
aid, (red,), (p1,) = new_auction('End during a sale', teams=(('Reds', 100),))
a, b = on(aid), on(aid)
def end_soon():
    time.sleep(0.05)   # just after the sale has started
    return b.post('/api/auction/end')
r = together(lambda: a.post('/api/sell_player', json={'player_id': p1, 'team_id': red, 'sold_price': 5}), end_soon)
with A.app.test_request_context():
    kept = json.loads(A.auction_row(aid)['report_json'])   # the record kept after the purge
in_report = [p['name'] for t in kept['teams'] for p in t.get('players') or []]
sold_now = status(p1)['status'] == 'sold'
check('a sale made as the auction ends is in the kept report', sold_now == ('P1' in in_report),
      'sale %s, report %s' % (r[0].status_code, in_report))
late = a.post('/api/sell_player', json={'player_id': p1, 'team_id': red, 'sold_price': 5})
check('nothing changes after the end', late.status_code == 403, str(late.status_code))
r = together(lambda: a.post('/api/auction/reopen'), lambda: b.post('/api/auction/reopen'))
r = together(lambda: a.post('/api/auction/end'), lambda: b.post('/api/auction/end'))
check('End Auction tapped twice at once ends it once', all(x.status_code == 200 for x in r)
      and sorted(bool(x.get_json().get('already_ended')) for x in r) == [False, True], str([x.get_json() for x in r]))

print('Two tabs, two auctions')
first, _, (fp,) = new_auction('First')
second, (s_red, _), (sp,) = new_auction('Second')
tab = admin_client()
tab.post('/api/auctions/%d/open' % first)
pinned = tab.get('/api/teams')
check('answers name their auction', pinned.headers.get('X-Auction-Id') == str(first))
tab.post('/api/auctions/%d/open' % second)   # what a second tab of this browser does
r = tab.post('/api/sell_player', json={'player_id': sp, 'team_id': s_red, 'sold_price': 1},
             headers={'X-Auction-Id': str(first)})
check("a sale from the first auction's tab is refused", r.status_code == 409 and r.get_json()['error'] == 'auction_switched',
      str(r.status_code))
check('and nothing in the other auction changed', all(p['status'] == 'unsold' for p in tab.get('/api/players').get_json()))
check('a request naming the open auction goes through',
      tab.get('/api/teams', headers={'X-Auction-Id': str(second)}).status_code == 200)

print('Background work')
aid, _, (p1,) = new_auction('Photo download', players=(('P1', 1),))
with A.app.test_request_context():
    A.g.auction_id = aid
    conn = A.get_db()
    conn.execute("UPDATE players SET photo_url = 'https://drive.google.com/file/d/abc/view' WHERE id = ?", (p1,))
    conn.commit()
def download_while_admin_uploads(url, auction_id, player_id):
    conn = A.open_auction_conn(auction_id)
    conn.execute("UPDATE players SET photo_url = '/media/admin.png' WHERE id = ?", (player_id,))
    conn.commit()
    conn.close()
    return '/media/from_drive.png'
_download = A.download_drive_photo
A.download_drive_photo = download_while_admin_uploads
A._hydrate_drive_photos(aid)
A.download_drive_photo = _download
check('a photo uploaded during a Drive download is kept', status(p1)['photo_url'] == '/media/admin.png',
      status(p1)['photo_url'])

aid, (red,), pids = new_auction('Backups', teams=(('Reds', 100),), players=[('P%d' % i, 1) for i in range(6)])
for pid in pids:
    admin.post('/api/sell_player', json={'player_id': pid, 'team_id': red, 'sold_price': 1})
errors = []
_print = A.print if hasattr(A, 'print') else print
import builtins
def capture(*a, **k):
    if a and 'Excel Backup' in str(a[0]): errors.append(str(a[0]))
    return _print(*a, **k)
builtins.print = capture
threads = [threading.Thread(target=A._write_excel_backup, args=(aid,)) for _ in range(12)]
for t in threads: t.start()
for t in threads: t.join()
builtins.print = _print
import openpyxl
wb = openpyxl.load_workbook(A.excel_backup_path(aid))
check('twelve backups at once: none fails', not errors, '; '.join(errors[:2]))
check('and the file is whole and current', wb['Sold Players'].max_row == 1 + len(pids), str(wb['Sold Players'].max_row))

print('\nALL CHECKS PASSED' if not check.failed else '\n%d CHECK(S) FAILED' % check.failed)
sys.exit(1 if check.failed else 0)
